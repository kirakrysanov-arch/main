"""Builds The Match Book stats tracker workbook. Run: python3 build_tracker.py"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "The-Match-Book-Stats-Tracker.xlsx"
GREEN, GREEN2, BRASS, CREAM, PAPER, INK, MUTED = "1D3A2A", "2D5440", "A07F42", "EFE7D3", "FBF8F0", "1F2A22", "7F8577"
FIRST, LAST = 6, 155  # Match Log data rows (150 matches)
SURFACES = ["Hard", "Clay", "Grass", "Indoor"]

title_font = Font(name="Georgia", size=22, color=GREEN)
sub_font = Font(name="Arial", size=9, color=BRASS, bold=True)
head_font = Font(name="Arial", size=9, bold=True, color="F3E7C4")
body_font = Font(name="Arial", size=10, color=INK)
calc_font = Font(name="Arial", size=10, color=GREEN2, bold=True)
note_font = Font(name="Arial", size=9, italic=True, color=MUTED)
head_fill = PatternFill("solid", fgColor=GREEN)
input_fill = PatternFill("solid", fgColor=PAPER)
calc_fill = PatternFill("solid", fgColor=CREAM)
hair = Side(style="thin", color="D9D1BA")
box = Border(left=hair, right=hair, top=hair, bottom=hair)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def title(ws, text, sub, width_cols):
    ws.sheet_view.showGridLines = False
    ws["B2"] = text
    ws["B2"].font = title_font
    ws["B3"] = sub.upper()
    ws["B3"].font = sub_font
    ws.row_dimensions[2].height = 34
    ws.column_dimensions["A"].width = 3
    for c in range(2, width_cols + 2):
        ws.cell(row=4, column=c).border = Border(bottom=Side(style="medium", color=BRASS))


def header_row(ws, row, col, labels, height=34):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, center, box
    ws.row_dimensions[row].height = height


def style_range(ws, rows, cols, fill, font, fmt=None, align=None):
    for r in rows:
        for c in cols:
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.font, cell.border = fill, font, box
            if fmt:
                cell.number_format = fmt
            if align:
                cell.alignment = align


wb = Workbook()

# ------------------------------------------------------------------ Start Here
ws = wb.active
ws.title = "Start Here"
title(ws, "The Match Book", "Tennis stats tracker · start here", 6)
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 90
rows = [
    ("How it works", "Log every match on the Match Log tab. The Dashboard, Opponents and monthly charts update by themselves."),
    ("Cream cells", "These are yours to fill in. Type or pick from the drop-down lists."),
    ("Beige cells", "These calculate automatically. Please don't type over them."),
    ("Sample data", "Rows 6–21 of the Match Log hold 16 sample matches so you can see the dashboard working. "
                    "Select the cream cells in those rows and press Delete to clear them before you start."),
    ("During a match", "Ask a parent, coach or friend to keep a tally on the Match Book chart page, "
                       "or note the totals shown by your club's scoreboard or a tracking app."),
    ("What to count", "1st serves in / attempted, points won on 1st and 2nd serve, break points won / chances, "
                      "winners, unforced errors and net points won / played. Leave anything you didn't track blank."),
    ("Google Sheets", "Upload this file to Google Drive, then open it with Google Sheets "
                      "(File › Save as Google Sheets). Everything works in Excel, Numbers and Google Sheets."),
    ("Opponents tab", "Type each opponent's name once (exactly as on the Match Log) to see your head-to-head record."),
    ("Season Plan tab", "List upcoming tournaments and entry deadlines. Days to go counts down automatically."),
]
for i, (k, v) in enumerate(rows):
    r = 6 + i * 2
    ws.cell(row=r, column=2, value=k).font = Font(name="Georgia", size=12, color=GREEN)
    c = ws.cell(row=r, column=3, value=v)
    c.font, c.alignment = body_font, Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
ws["B6"].fill = input_fill
ws["B8"].fill = input_fill
ws["B10"].fill = calc_fill
ws["B26"] = "For personal use only. Thank you for supporting a small studio."
ws["B26"].font = note_font

# ------------------------------------------------------------------ Match Log
ml = wb.create_sheet("Match Log")
inputs = ["Date", "Event", "Round", "Opponent", "Surface", "Result\n(W / L)", "Score", "Sets\nwon", "Sets\nlost",
          "Aces", "Double\nfaults", "1st serves\nin", "1st serves\nattempted", "1st serve\npts won", "2nd serve\npts won",
          "2nd serve\npts played", "Break pts\nwon", "Break pt\nchances", "Winners", "Unforced\nerrors",
          "Net pts\nwon", "Net pts\nplayed"]
calcs = ["1st serve\n%", "1st serve\npts won %", "2nd serve\npts won %", "Break pt\nconversion", "Winners −\nerrors",
         "Net pts\nwon %", "Month"]
title(ml, "Match Log", "One row per match · cream cells are yours to fill in", len(inputs) + len(calcs))
header_row(ml, 5, 2, inputs + calcs, 40)
widths = [12, 22, 9, 18, 10, 9, 16] + [8] * 15 + [10] * 7
for i, w in enumerate(widths):
    ml.column_dimensions[get_column_letter(2 + i)].width = w
ml.freeze_panes = "C6"

ni = len(inputs)
col = {name: get_column_letter(2 + i) for i, name in enumerate(
    ["date", "event", "round", "opp", "surface", "result", "score", "sw", "sl", "aces", "df", "fin", "fatt", "fwon",
     "swon", "splay", "bpw", "bpc", "win", "ue", "netw", "netp",
     "c_fs", "c_fsw", "c_ssw", "c_bp", "c_wue", "c_net", "c_month"])}
style_range(ml, range(FIRST, LAST + 1), range(2, 2 + ni), input_fill, body_font, align=Alignment(horizontal="center"))
style_range(ml, range(FIRST, LAST + 1), range(2 + ni, 2 + ni + len(calcs)), calc_fill, calc_font, "0%",
            Alignment(horizontal="center"))
for r in range(FIRST, LAST + 1):
    ml[f"{col['date']}{r}"].number_format = "d mmm yyyy"
    for k in ("event", "opp", "score"):
        ml[f"{col[k]}{r}"].alignment = Alignment(horizontal="left")
    c = col
    ml[f"{c['c_fs']}{r}"] = f'=IF(OR({c["fatt"]}{r}="",{c["fatt"]}{r}=0),"",{c["fin"]}{r}/{c["fatt"]}{r})'
    ml[f"{c['c_fsw']}{r}"] = f'=IF(OR({c["fin"]}{r}="",{c["fin"]}{r}=0),"",{c["fwon"]}{r}/{c["fin"]}{r})'
    ml[f"{c['c_ssw']}{r}"] = f'=IF(OR({c["splay"]}{r}="",{c["splay"]}{r}=0),"",{c["swon"]}{r}/{c["splay"]}{r})'
    ml[f"{c['c_bp']}{r}"] = f'=IF(OR({c["bpc"]}{r}="",{c["bpc"]}{r}=0),"",{c["bpw"]}{r}/{c["bpc"]}{r})'
    ml[f"{c['c_wue']}{r}"] = f'=IF(OR({c["win"]}{r}="",{c["ue"]}{r}=""),"",{c["win"]}{r}-{c["ue"]}{r})'
    ml[f"{c['c_wue']}{r}"].number_format = "+0;-0;0"
    ml[f"{c['c_net']}{r}"] = f'=IF(OR({c["netp"]}{r}="",{c["netp"]}{r}=0),"",{c["netw"]}{r}/{c["netp"]}{r})'
    ml[f"{c['c_month']}{r}"] = f'=IF({c["date"]}{r}="","",MONTH({c["date"]}{r}))'
    ml[f"{c['c_month']}{r}"].number_format = "0"

dv_surface = DataValidation(type="list", formula1='"Hard,Clay,Grass,Indoor"', allow_blank=True)
dv_result = DataValidation(type="list", formula1='"W,L"', allow_blank=True)
ml.add_data_validation(dv_surface)
ml.add_data_validation(dv_result)
dv_surface.add(f"{col['surface']}{FIRST}:{col['surface']}{LAST}")
dv_result.add(f"{col['result']}{FIRST}:{col['result']}{LAST}")
rng = f"{col['result']}{FIRST}:{col['result']}{LAST}"
ml.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"W"'], font=Font(color="2D6A3E", bold=True)))
ml.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"L"'], font=Font(color="9A3B2E", bold=True)))

# Sample matches: date, event, round, opponent, surface, result, score, sets won, sets lost, then the counted stats.
samples = [
    (date(2026, 3, 7), "Spring Open", "R1", "J. Carter", "Hard", "W", "6-3 6-4", 2, 0, 4, 3, 38, 62, 27, 13, 24, 4, 9, 18, 21, 8, 12),
    (date(2026, 3, 8), "Spring Open", "R2", "M. Alvarez", "Hard", "L", "4-6 6-7", 0, 2, 2, 5, 35, 66, 22, 12, 31, 2, 7, 14, 29, 5, 10),
    (date(2026, 3, 21), "Club League", "RR", "S. Patel", "Hard", "W", "6-2 6-1", 2, 0, 5, 1, 34, 50, 26, 11, 16, 5, 8, 20, 14, 9, 11),
    (date(2026, 4, 11), "Clay Classic", "R1", "L. Novak", "Clay", "W", "7-5 3-6 6-4", 2, 1, 1, 4, 52, 84, 33, 17, 32, 6, 14, 22, 33, 6, 11),
    (date(2026, 4, 12), "Clay Classic", "QF", "R. Moreau", "Clay", "L", "2-6 5-7", 0, 2, 1, 6, 36, 63, 20, 11, 27, 1, 6, 11, 31, 3, 7),
    (date(2026, 4, 25), "Club League", "RR", "J. Carter", "Hard", "W", "6-4 6-4", 2, 0, 3, 2, 40, 64, 29, 14, 24, 3, 7, 19, 20, 7, 10),
    (date(2026, 5, 16), "County Champs", "R1", "D. Okafor", "Clay", "W", "6-1 6-3", 2, 0, 3, 2, 36, 52, 27, 10, 16, 5, 9, 17, 15, 6, 8),
    (date(2026, 5, 17), "County Champs", "R2", "M. Alvarez", "Clay", "W", "6-7 6-4 7-6", 2, 1, 5, 4, 58, 91, 39, 19, 33, 4, 12, 26, 30, 11, 17),
    (date(2026, 5, 18), "County Champs", "SF", "A. Lindqvist", "Clay", "L", "3-6 4-6", 0, 2, 2, 3, 37, 60, 23, 11, 23, 2, 8, 13, 26, 4, 8),
    (date(2026, 6, 13), "Grass Invitational", "R1", "T. Becker", "Grass", "W", "7-6 6-4", 2, 0, 9, 3, 41, 63, 34, 12, 22, 2, 5, 27, 18, 14, 19),
    (date(2026, 6, 14), "Grass Invitational", "QF", "S. Patel", "Grass", "W", "6-3 7-5", 2, 0, 7, 2, 40, 58, 32, 11, 18, 3, 6, 24, 16, 12, 16),
    (date(2026, 6, 15), "Grass Invitational", "SF", "R. Moreau", "Grass", "L", "6-7 4-6", 0, 2, 6, 4, 39, 65, 30, 12, 26, 0, 4, 21, 22, 10, 17),
    (date(2026, 7, 11), "Summer Series", "R1", "L. Novak", "Hard", "W", "6-2 6-2", 2, 0, 6, 1, 35, 51, 28, 11, 16, 5, 8, 21, 12, 8, 10),
    (date(2026, 8, 22), "Summer Series", "QF", "A. Lindqvist", "Hard", "W", "4-6 6-3 6-3", 2, 1, 4, 3, 50, 78, 35, 16, 28, 5, 11, 25, 24, 9, 13),
    (date(2026, 9, 19), "Club Championship", "SF", "D. Okafor", "Indoor", "W", "6-4 6-2", 2, 0, 5, 2, 37, 55, 28, 11, 18, 4, 8, 22, 15, 8, 11),
    (date(2026, 10, 10), "Club Championship", "F", "M. Alvarez", "Indoor", "W", "7-5 6-4", 2, 0, 6, 2, 42, 63, 32, 13, 21, 3, 7, 24, 19, 9, 13),
]
for i, s in enumerate(samples):
    for j, v in enumerate(s):
        ml.cell(row=FIRST + i, column=2 + j, value=v)

# ------------------------------------------------------------------ Dashboard
db = wb.create_sheet("Dashboard", 1)
title(db, "Season Dashboard", "Updates automatically from the Match Log", 10)
for c, w in zip("BCDEFGHIJK", [24, 13, 4, 24, 13, 4, 24, 13, 4, 4]):
    db.column_dimensions[c].width = w
R = lambda k: f"'Match Log'!${col[k]}${FIRST}:${col[k]}${LAST}"


def safe_div(num, den):
    return f"=IF({den}=0,0,{num}/{den})"


kpis = [
    [("Matches played", f"=COUNT({R('date')})", "0"),
     ("Wins", f'=COUNTIF({R("result")},"W")', "0"),
     ("Losses", f'=COUNTIF({R("result")},"L")', "0")],
    [("Win rate", "=IF(C7=0,0,F7/C7)", "0%"),
     ("Sets won", f"=SUM({R('sw')})", "0"),
     ("Sets lost", f"=SUM({R('sl')})", "0")],
    [("1st serve in", safe_div(f"SUM({R('fin')})", f"SUM({R('fatt')})"), "0%"),
     ("1st serve pts won", safe_div(f"SUM({R('fwon')})", f"SUM({R('fin')})"), "0%"),
     ("2nd serve pts won", safe_div(f"SUM({R('swon')})", f"SUM({R('splay')})"), "0%")],
    [("Break pt conversion", safe_div(f"SUM({R('bpw')})", f"SUM({R('bpc')})"), "0%"),
     ("Net pts won", safe_div(f"SUM({R('netw')})", f"SUM({R('netp')})"), "0%"),
     ("Winners ÷ errors", safe_div(f"SUM({R('win')})", f"SUM({R('ue')})"), "0.00"),
     ],
    [("Aces per match", safe_div(f"SUM({R('aces')})", f"COUNT({R('aces')})"), "0.0"),
     ("Double faults per match", safe_div(f"SUM({R('df')})", f"COUNT({R('df')})"), "0.0"),
     ("Unforced errors per match", safe_div(f"SUM({R('ue')})", f"COUNT({R('ue')})"), "0.0")],
]
for i, rowk in enumerate(kpis):
    r = 6 + i * 3
    for j, (label, f, fmt) in enumerate(rowk):
        lc, vc = 2 + j * 3, 3 + j * 3
        lab = db.cell(row=r, column=lc, value=label.upper())
        lab.font = Font(name="Arial", size=8, bold=True, color=BRASS)
        db.merge_cells(start_row=r + 1, start_column=lc, end_row=r + 1, end_column=vc)
        v = db.cell(row=r + 1, column=lc, value=f.replace("C7", "B7").replace("F7", "E7"))
        v.font, v.number_format, v.alignment = Font(name="Georgia", size=24, color=GREEN), fmt, Alignment(horizontal="left")
        db.row_dimensions[r + 1].height = 36
        for cc in (lc, vc):
            db.cell(row=r + 1, column=cc).border = Border(bottom=hair)

# Surface breakdown
r0 = 22
db.cell(row=r0 - 1, column=2, value="BY SURFACE").font = sub_font
header_row(db, r0, 2, ["Surface", "Played", "", "Won", "Win rate", "", "1st serve in", "BP conversion"], 24)
for i, s in enumerate(SURFACES):
    r = r0 + 1 + i
    db.cell(row=r, column=2, value=s)
    db.cell(row=r, column=3, value=f'=COUNTIF({R("surface")},B{r})')
    db.cell(row=r, column=5, value=f'=COUNTIFS({R("surface")},B{r},{R("result")},"W")')
    db.cell(row=r, column=6, value=f"=IF(C{r}=0,0,E{r}/C{r})")
    db.cell(row=r, column=8, value=f'=IF(SUMIF({R("surface")},B{r},{R("fatt")})=0,0,SUMIF({R("surface")},B{r},{R("fin")})/SUMIF({R("surface")},B{r},{R("fatt")}))')
    db.cell(row=r, column=9, value=f'=IF(SUMIF({R("surface")},B{r},{R("bpc")})=0,0,SUMIF({R("surface")},B{r},{R("bpw")})/SUMIF({R("surface")},B{r},{R("bpc")}))')
    for c in range(2, 10):
        cell = db.cell(row=r, column=c)
        cell.font, cell.border, cell.fill = body_font, box, calc_fill
        cell.alignment = Alignment(horizontal="center")
    for c in (6, 8, 9):
        db.cell(row=r, column=c).number_format = "0%"
    db.cell(row=r, column=2).alignment = Alignment(horizontal="left")
db.column_dimensions["I"].width = 14

# Monthly table (feeds the charts)
m0 = 29
db.cell(row=m0 - 1, column=2, value="BY MONTH").font = sub_font
header_row(db, m0, 2, ["Month", "Wins", "", "Losses", "Avg winners", "", "Avg errors", "1st serve in"], 24)
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for i, m in enumerate(MONTHS):
    r = m0 + 1 + i
    mc = R("c_month")
    db.cell(row=r, column=2, value=m)
    db.cell(row=r, column=3, value=f'=COUNTIFS({mc},{i + 1},{R("result")},"W")')
    db.cell(row=r, column=5, value=f'=COUNTIFS({mc},{i + 1},{R("result")},"L")')
    db.cell(row=r, column=6, value=f'=IF(COUNTIFS({mc},{i + 1},{R("win")},"<>")=0,0,SUMIF({mc},{i + 1},{R("win")})/COUNTIFS({mc},{i + 1},{R("win")},"<>"))')
    db.cell(row=r, column=8, value=f'=IF(COUNTIFS({mc},{i + 1},{R("ue")},"<>")=0,0,SUMIF({mc},{i + 1},{R("ue")})/COUNTIFS({mc},{i + 1},{R("ue")},"<>"))')
    db.cell(row=r, column=9, value=f'=IF(SUMIF({mc},{i + 1},{R("fatt")})=0,0,SUMIF({mc},{i + 1},{R("fin")})/SUMIF({mc},{i + 1},{R("fatt")}))')
    for c in range(2, 10):
        cell = db.cell(row=r, column=c)
        cell.font, cell.border, cell.fill = body_font, box, calc_fill
        cell.alignment = Alignment(horizontal="center")
    db.cell(row=r, column=6).number_format = "0.0"
    db.cell(row=r, column=8).number_format = "0.0"
    db.cell(row=r, column=9).number_format = "0%"
# Blank zero cells keep the tables calm when a month has no matches.
db.conditional_formatting.add(f"C{m0 + 1}:I{m0 + 12}", CellIsRule(operator="equal", formula=["0"], font=Font(color="C9C1A8")))


def bar(title_text, cols, colors, anchor, stacked=False):
    ch = BarChart()
    ch.type = "col"
    ch.title = title_text
    ch.style = 2
    ch.height, ch.width = 7.2, 15.5
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
    for c, color in zip(cols, colors):
        ch.add_data(Reference(db, min_col=c, min_row=m0, max_row=m0 + 12), titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = color
        ch.series[-1].graphicalProperties.line.solidFill = color
    ch.set_categories(Reference(db, min_col=2, min_row=m0 + 1, max_row=m0 + 12))
    ch.y_axis.majorGridlines = None
    ch.legend.position = "b"
    db.add_chart(ch, anchor)


bar("Wins and losses by month", [3, 5], [GREEN, "C8B27A"], "L6", stacked=True)
bar("Winners vs unforced errors", [6, 8], [GREEN2, "B5654F"], "L22")

# ------------------------------------------------------------------ Opponents
op = wb.create_sheet("Opponents")
title(op, "Head to Head", "Type each opponent once, spelled exactly as in the Match Log", 6)
header_row(op, 5, 2, ["Opponent", "Played", "Won", "Lost", "Win rate", "Last played"], 28)
for c, w in zip("BCDEFG", [24, 10, 10, 10, 11, 15]):
    op.column_dimensions[c].width = w
names = sorted({s[3] for s in samples})
for i in range(40):
    r = 6 + i
    if i < len(names):
        op.cell(row=r, column=2, value=names[i])
    op.cell(row=r, column=3, value=f'=IF(B{r}="","",COUNTIF({R("opp")},B{r}))')
    op.cell(row=r, column=4, value=f'=IF(B{r}="","",COUNTIFS({R("opp")},B{r},{R("result")},"W"))')
    op.cell(row=r, column=5, value=f'=IF(B{r}="","",COUNTIFS({R("opp")},B{r},{R("result")},"L"))')
    op.cell(row=r, column=6, value=f'=IF(OR(B{r}="",C{r}=0),"",D{r}/C{r})')
    op.cell(row=r, column=7, value=f'=IF(OR(B{r}="",C{r}=0),"",_xlfn.MAXIFS({R("date")},{R("opp")},B{r}))')
    style_range(op, [r], [2], input_fill, body_font)
    style_range(op, [r], range(3, 8), calc_fill, calc_font, align=Alignment(horizontal="center"))
    op.cell(row=r, column=6).number_format = "0%"
    op.cell(row=r, column=7).number_format = "d mmm yyyy"
op.freeze_panes = "B6"

# ------------------------------------------------------------------ Season Plan
sp = wb.create_sheet("Season Plan")
title(sp, "Season Plan", "Upcoming tournaments and deadlines", 8)
header_row(sp, 5, 2, ["Start date", "Tournament / league", "Location", "Surface", "Entry deadline", "Entered?",
                      "Days to go", "Goal"], 28)
for c, w in zip("BCDEFGHI", [13, 28, 18, 10, 14, 10, 11, 28]):
    sp.column_dimensions[c].width = w
dv_s2 = DataValidation(type="list", formula1='"Hard,Clay,Grass,Indoor"', allow_blank=True)
dv_yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
sp.add_data_validation(dv_s2)
sp.add_data_validation(dv_yn)
dv_s2.add("E6:E45")
dv_yn.add("G6:G45")
example = (date(2027, 3, 6), "Spring Open", "City Tennis Centre", "Hard", date(2027, 2, 20), "Yes", None, "Reach the quarter-finals")
for i in range(40):
    r = 6 + i
    style_range(sp, [r], [2, 3, 4, 5, 6, 7, 9], input_fill, body_font)
    sp.cell(row=r, column=8, value=f'=IF(B{r}="","",MAX(0,B{r}-TODAY()))')
    style_range(sp, [r], [8], calc_fill, calc_font, "0", Alignment(horizontal="center"))
    for c in (2, 6):
        sp.cell(row=r, column=c).number_format = "d mmm yyyy"
for j, v in enumerate(example):
    if v is not None:
        sp.cell(row=6, column=2 + j, value=v)
sp.freeze_panes = "B6"

for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = GREEN if sheet.title in ("Dashboard", "Match Log") else BRASS
wb.active = 0
wb.save(OUT)
print(OUT)
