"""Builds the Tennis Round Robin & Ladder Organizer workbook. Run: python3 build_organizer.py [output.xlsx]"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else (
    Path(__file__).resolve().parent.parent / "listing" / "upload" / "Tennis-Round-Robin-and-Ladder-Organizer.xlsx")

# Retro Club palette
GREEN, NAVY, RED, CREAM, PAPER, MUTED, HAIR = "1F6B3B", "1B2B4B", "C2412D", "EBDDBC", "FBF5E8", "8C8068", "D8C9A5"
MAX_PLAYERS = 16          # round robin
SLOTS = MAX_PLAYERS // 2  # matches per round at most
SCHED_ROWS = (MAX_PLAYERS - 1) * SLOTS  # 120 rows covers every round for up to 16 players
LADDER_MAX = 20
CHALLENGES = 100

title_font = Font(name="Georgia", size=22, bold=True, color=NAVY)
sub_font = Font(name="Arial", size=9, bold=True, color=RED)
head_font = Font(name="Arial", size=9, bold=True, color="F6EDD9")
label_font = Font(name="Arial", size=10, bold=True, color=NAVY)
body_font = Font(name="Arial", size=10, color=NAVY)
calc_font = Font(name="Arial", size=10, bold=True, color=GREEN)
note_font = Font(name="Arial", size=9, italic=True, color=MUTED)
head_fill = PatternFill("solid", fgColor=GREEN)
input_fill = PatternFill("solid", fgColor=PAPER)
calc_fill = PatternFill("solid", fgColor=CREAM)
band_fill = PatternFill("solid", fgColor="F3E9D2")
hair = Side(style="thin", color=HAIR)
box = Border(left=hair, right=hair, top=hair, bottom=hair)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center")


def title(ws, text, sub, last_col):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws["B2"], ws["B3"] = text, sub.upper()
    ws["B2"].font, ws["B3"].font = title_font, sub_font
    ws.row_dimensions[2].height = 34
    for c in range(2, last_col + 1):
        ws.cell(row=4, column=c).border = Border(bottom=Side(style="medium", color=RED))


def header(ws, row, first_col, labels, height=30):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=first_col + i, value=label)
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, center, box
    ws.row_dimensions[row].height = height


def cell(ws, ref, value=None, kind="calc", fmt=None, align=center):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.fill = {"input": input_fill, "calc": calc_fill}[kind]
    c.font = {"input": body_font, "calc": calc_font}[kind]
    c.border, c.alignment = box, align
    if fmt:
        c.number_format = fmt
    return c


wb = Workbook()

# ------------------------------------------------------------------ Start Here
ws = wb.active
ws.title = "Start Here"
title(ws, "Round Robin & Ladder Organizer", "Start here", 3)
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 92
guide = [
    ("Cream cells", "Yours to fill in. Type, or pick from the drop-down lists.", "input"),
    ("Beige cells", "Calculate by themselves. Please don't type over them.", "calc"),
    ("1 · Setup", "Name your event, set the number of courts and the points for a win or tie, then list up to 16 players "
                  "or doubles teams from the top, with no gaps.", None),
    ("2 · Schedule", "The full schedule appears instantly: every player meets every other player once, with court numbers. "
                     "With an odd number of players, one player rests (a bye) each round.", None),
    ("3 · Enter scores", "Type the games won by each side, for example 6 and 4. The winner fills in by itself.", None),
    ("4 · Standings", "Ranked by points, then game difference, then games won. The Results Grid shows every score at a glance.", None),
    ("Ladder", "For a season-long challenge ladder: list up to 20 players in starting order on the Ladder tab, then log "
               "each challenge on the Challenges tab. A challenger who beats a higher-ranked player takes their place, "
               "and everyone in between moves down one.", None),
    ("Sample data", "The sample players and scores show how everything works. Before your event, clear the cream cells "
                    "on Setup, the score columns on Schedule, and the Ladder and Challenges tabs.", None),
    ("Changing players", "If you change the player list after scores are entered, clear the score columns first: "
                         "the schedule is rebuilt for the new list.", None),
    ("Google Sheets", "Upload this file to Google Drive, open it, then choose File › Save as Google Sheets. "
                      "It also works in Microsoft Excel and Apple Numbers.", None),
    ("Printing", "Each tab is set up to print on one page width. Print the Schedule for the noticeboard, "
                 "or use the blank Draw Sheets PDF.", None),
]
for i, (k, v, kind) in enumerate(guide):
    r = 6 + i * 2
    a = ws.cell(row=r, column=2, value=k)
    a.font = Font(name="Georgia", size=12, bold=True, color=GREEN)
    if kind:
        a.fill = input_fill if kind == "input" else calc_fill
    b = ws.cell(row=r, column=3, value=v)
    b.font, b.alignment = body_font, Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
ws.cell(row=6 + len(guide) * 2, column=2, value="For personal and single-club use. Thank you for supporting a small studio.").font = note_font

# ------------------------------------------------------------------ Setup
su = wb.create_sheet("Setup")
title(su, "Setup", "Event details and players", 4)
su.column_dimensions["B"].width = 26
su.column_dimensions["C"].width = 34
su.column_dimensions["D"].width = 4
su.column_dimensions["E"].width = 50
fields = [
    (6, "Event name", "Saturday Social Round Robin", "input", None),
    (7, "Format", "Singles", "input", None),
    (8, "Courts available", 3, "input", "0"),
    (9, "Points for a win", 1, "input", "0.0"),
    (10, "Points for a tie", 0.5, "input", "0.0"),
    (11, "Match format (note)", "One set, first to 6 games", "input", None),
    (12, "Players entered", "=COUNTA($C$16:$C$31)", "calc", "0"),
    (13, "Rounds", "=IF($C$12<3,0,$C$12+MOD($C$12,2)-1)", "calc", "0"),
]
for r, label, value, kind, fmt in fields:
    su.cell(row=r, column=2, value=label).font = label_font
    cell(su, f"C{r}", value, kind, fmt, left)
su["E8"] = "Matches beyond the number of courts show “Next free”."
su["E12"] = "Enter at least 3. The schedule supports up to 16."
su["E13"] = "Every player meets every other player once."
for r in (8, 12, 13):
    su[f"E{r}"].font = note_font
dv_format = DataValidation(type="list", formula1='"Singles,Doubles"', allow_blank=True)
su.add_data_validation(dv_format)
dv_format.add("C7")
header(su, 15, 2, ["No.", "Player or doubles team"], 24)
sample_players = ["Anna", "Ben", "Carla", "David", "Elena", "Felix", "Grace"]
for i in range(MAX_PLAYERS):
    r = 16 + i
    n = su.cell(row=r, column=2, value=i + 1)
    n.font, n.alignment, n.border = label_font, center, box
    cell(su, f"C{r}", sample_players[i] if i < len(sample_players) else None, "input", align=left)
su["E16"] = "For doubles, enter each team as one line, e.g. “Anna & Ben”."
su["E16"].font = note_font

# ------------------------------------------------------------------ Schedule
sc = wb.create_sheet("Schedule")
title(sc, "Schedule", "Created automatically · enter games won in the cream cells", 8)
header(sc, 5, 2, ["Round", "Court", "Player / team A", "Games A", "Games B", "Player / team B", "Winner"], 30)
for c, w in zip("BCDEFGH", [9, 11, 26, 10, 10, 26, 24]):
    sc.column_dimensions[c].width = w
for c in "JKLMNOPQR":
    sc.column_dimensions[c].hidden = True
FIRST, LAST = 6, 6 + SCHED_ROWS - 1
NAMES = "Setup!$C$16:$C$31"
N, ROUNDS, COURTS = "Setup!$C$12", "Setup!$C$13", "Setup!$C$8"
for k in range(SCHED_ROWS):
    r = FIRST + k
    # Row k holds match slot m of round r, packed with no gaps: each round has (rounds + 1) / 2 slots.
    sc[f"J{r}"] = f'=IF({ROUNDS}=0,"",INT({k}/(({ROUNDS}+1)/2)))'
    sc[f"K{r}"] = f'=IF({ROUNDS}=0,"",MOD({k},({ROUNDS}+1)/2))'
    sc[f"L{r}"] = f"=AND({ROUNDS}>0,J{r}<{ROUNDS})"   # is this row a real match?
    # Circle method: the last seat is fixed and meets the round's player in slot 0; other seats pair r+m with r-m.
    sc[f"M{r}"] = f'=IF(NOT(L{r}),"",IF(K{r}=0,J{r},MOD(J{r}+K{r},{ROUNDS})))'
    sc[f"N{r}"] = f'=IF(NOT(L{r}),"",IF(K{r}=0,{ROUNDS},MOD(J{r}-K{r}+{ROUNDS},{ROUNDS})))'
    sc[f"O{r}"] = f'=AND(L{r},OR(M{r}>={N},N{r}>={N}))'       # bye: the empty seat when players are odd
    sc[f"P{r}"] = f'=IF(AND(L{r},NOT(O{r}),E{r}<>"",F{r}<>""),1,0)'  # result entered
    sc[f"Q{r}"] = f'=IF(L{r},D{r}&"|"&G{r},"")'
    sc[f"R{r}"] = f'=IF(L{r},G{r}&"|"&D{r},"")'
    cell(sc, f"B{r}", f'=IF(L{r},J{r}+1,"")')
    court = f"IF(MOD({N},2)=1,K{r},K{r}+1)"
    cell(sc, f"C{r}", f'=IF(NOT(L{r}),"",IF(O{r},"Bye",IF({court}>{COURTS},"Next free",{court})))')
    cell(sc, f"D{r}", f'=IF(NOT(L{r}),"",IF(M{r}>={N},"Bye",INDEX({NAMES},M{r}+1)))', align=left)
    cell(sc, f"E{r}", None, "input", "0")
    cell(sc, f"F{r}", None, "input", "0")
    cell(sc, f"G{r}", f'=IF(NOT(L{r}),"",IF(N{r}>={N},"Bye",INDEX({NAMES},N{r}+1)))', align=left)
    cell(sc, f"H{r}", f'=IF(OR(NOT(L{r}),O{r},E{r}="",F{r}=""),"",IF(E{r}>F{r},D{r},IF(F{r}>E{r},G{r},"Tie")))', align=left)
rng = f"B{FIRST}:H{LAST}"
sc.conditional_formatting.add(rng, FormulaRule(formula=[f"$O{FIRST}"], font=Font(italic=True, color=MUTED)))
sc.conditional_formatting.add(rng, FormulaRule(formula=[f"AND($L{FIRST},MOD($J{FIRST},2)=1)"], fill=band_fill))
sc.conditional_formatting.add(rng, FormulaRule(formula=[f"NOT($L{FIRST})"], fill=PatternFill("solid", fgColor="FFFFFF"),
                                                border=Border()))
dv_games = DataValidation(type="whole", operator="between", formula1="0", formula2="99", allow_blank=True)
sc.add_data_validation(dv_games)
dv_games.add(f"E{FIRST}:F{LAST}")
sc.freeze_panes = "B6"
sc.print_title_rows = "5:5"
sc.print_area = f"B2:H{LAST}"

# Sample scores for the first three rounds of the 7-player sample.
sample_scores = [(6, 3), (4, 6), (6, 2), (6, 4), (5, 6), (6, 1), (3, 6), (5, 5), (6, 0)]
rows_with_matches = [FIRST + rnd * 4 + m for rnd in range(3) for m in (1, 2, 3)]  # 7 players: 4 slots a round
for r, (a, b) in zip(rows_with_matches, sample_scores):
    sc[f"E{r}"], sc[f"F{r}"] = a, b

# ------------------------------------------------------------------ Standings
st = wb.create_sheet("Standings")
title(st, "Standings", "Points, then game difference, then games won", 11)
header(st, 6, 2, ["Pos", "Player / team", "Played", "Won", "Tied", "Lost", "Games for", "Games against", "Game diff", "Points"], 30)
for c, w in zip("BCDEFGHIJK", [7, 26, 9, 8, 8, 8, 10, 10, 10, 10]):
    st.column_dimensions[c].width = w
for c in "MNOPQRSTUVW":
    st.column_dimensions[c].hidden = True
S = lambda c: f"Schedule!${c}${FIRST}:${c}${LAST}"
for i in range(MAX_PLAYERS):
    r = 7 + i
    p = f"Setup!$C${16 + i}"
    st[f"M{r}"] = f'=IF({p}="","",{p})'
    st[f"N{r}"] = f'=IF(M{r}="","",SUMIF({S("D")},M{r},{S("P")})+SUMIF({S("G")},M{r},{S("P")}))'
    st[f"O{r}"] = f'=IF(M{r}="","",COUNTIF({S("H")},M{r}))'
    st[f"P{r}"] = f'=IF(M{r}="","",COUNTIFS({S("D")},M{r},{S("H")},"Tie")+COUNTIFS({S("G")},M{r},{S("H")},"Tie"))'
    st[f"Q{r}"] = f'=IF(M{r}="","",N{r}-O{r}-P{r})'
    st[f"R{r}"] = f'=IF(M{r}="","",SUMIFS({S("E")},{S("D")},M{r},{S("P")},1)+SUMIFS({S("F")},{S("G")},M{r},{S("P")},1))'
    st[f"S{r}"] = f'=IF(M{r}="","",SUMIFS({S("F")},{S("D")},M{r},{S("P")},1)+SUMIFS({S("E")},{S("G")},M{r},{S("P")},1))'
    st[f"T{r}"] = f'=IF(M{r}="","",R{r}-S{r})'
    st[f"U{r}"] = f'=IF(M{r}="","",O{r}*Setup!$C$9+P{r}*Setup!$C$10)'
    # Sort key: points, then game difference, then games won; the last term keeps equal keys unique.
    st[f"V{r}"] = f'=IF(M{r}="","",U{r}*1000000+(T{r}+1000)*1000+R{r}+(100-ROW())/1000)'
    st[f"W{r}"] = f'=IF(ROW()-6>COUNT($V$7:$V$22),"",MATCH(LARGE($V$7:$V$22,ROW()-6),$V$7:$V$22,0))'
    cell(st, f"B{r}", f'=IF(W{r}="","",ROW()-6)')
    cell(st, f"C{r}", f'=IF(W{r}="","",INDEX($M$7:$M$22,W{r}))', align=left)
    for out_c, src in zip("DEFGHIJK", "NOPQRSTU"):
        cell(st, f"{out_c}{r}", f'=IF(W{r}="","",INDEX(${src}$7:${src}$22,W{r}))',
             fmt="+0;-0;0" if out_c == "J" else ("0.0" if out_c == "K" else "0"))
st.conditional_formatting.add("B7:K9", FormulaRule(formula=['$B7<>""'], font=Font(bold=True, color=RED)))
st.print_area = "B2:K22"

# ------------------------------------------------------------------ Results Grid
rg = wb.create_sheet("Results Grid")
title(rg, "Results Grid", "Read across: the row player's games first", 2 + MAX_PLAYERS)
rg.column_dimensions["B"].width = 20
rg.row_dimensions[7].height = 64
corner = rg["B7"]
corner.value, corner.font, corner.fill, corner.alignment = "vs", head_font, head_fill, center
for j in range(MAX_PLAYERS):
    c = col(3 + j)
    rg.column_dimensions[c].width = 9
    h = rg[f"{c}7"]
    h.value = f'=IF(Setup!$C${16 + j}="","",Setup!$C${16 + j})'
    h.font, h.fill, h.border = head_font, head_fill, box
    h.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True)
for i in range(MAX_PLAYERS):
    r = 8 + i
    rh = rg[f"B{r}"]
    rh.value = f'=IF(Setup!$C${16 + i}="","",Setup!$C${16 + i})'
    rh.font, rh.fill, rh.border, rh.alignment = Font(name="Arial", size=10, bold=True, color="F6EDD9"), head_fill, box, left
    for j in range(MAX_PLAYERS):
        c = col(3 + j)
        a, b = f"$B{r}", f"{c}$7"
        x = f'MATCH({a}&"|"&{b},{S("Q")},0)'
        y = f'MATCH({a}&"|"&{b},{S("R")},0)'
        f = (f'=IF(OR({a}="",{b}=""),"",IF({a}={b},"—",'
             f'IFERROR(IF(INDEX({S("P")},{x})=1,INDEX({S("E")},{x})&"–"&INDEX({S("F")},{x}),""),'
             f'IFERROR(IF(INDEX({S("P")},{y})=1,INDEX({S("F")},{y})&"–"&INDEX({S("E")},{y}),""),""))))')
        cell(rg, f"{c}{r}", f)
rg.conditional_formatting.add(f"C8:{col(2 + MAX_PLAYERS)}{7 + MAX_PLAYERS}",
                              FormulaRule(formula=['AND(C8<>"",C8<>"—",VALUE(LEFT(C8,FIND("–",C8)-1))>VALUE(MID(C8,FIND("–",C8)+1,3)))'],
                                          font=Font(bold=True, color=GREEN)))
rg.page_setup.orientation = "landscape"

# ------------------------------------------------------------------ Ladder
ld = wb.create_sheet("Ladder")
title(ld, "Challenge Ladder", "Starting order on the left · live ladder on the right", 10)
header(ld, 6, 2, ["Start", "Player (starting order)"], 30)
header(ld, 6, 5, ["Rank", "Player", "Challenges won", "Challenges lost", "Moved"], 30)
for c, w in zip("BCDEFGHI", [8, 26, 4, 8, 26, 12, 12, 10]):
    ld.column_dimensions[c].width = w
ENGINE_LAST = col(3 + CHALLENGES)  # engine column holding the position after the last challenge
sample_ladder = ["Olivia", "Marcus", "Priya", "Tom", "Sofia", "Jonas", "Leah", "Mateo", "Nina", "Oscar"]
CH = lambda c: f"Challenges!${c}$7:${c}${6 + CHALLENGES}"
for i in range(LADDER_MAX):
    r = 7 + i
    n = ld.cell(row=r, column=2, value=i + 1)
    n.font, n.alignment, n.border = label_font, center, box
    cell(ld, f"C{r}", sample_ladder[i] if i < len(sample_ladder) else None, "input", align=left)
    cell(ld, f"E{r}", f'=IF(F{r}="","",{i + 1})')
    cell(ld, f"F{r}", f"=IFERROR(INDEX(Engine!$B$7:$B$26,MATCH({i + 1},Engine!${ENGINE_LAST}$7:${ENGINE_LAST}$26,0)),\"\")",
         align=left)
    cell(ld, f"G{r}", f'=IF(F{r}="","",COUNTIFS({CH("D")},F{r},{CH("F")},"Challenger")+COUNTIFS({CH("E")},F{r},{CH("F")},"Defender"))', fmt="0")
    cell(ld, f"H{r}", f'=IF(F{r}="","",COUNTIFS({CH("D")},F{r},{CH("F")},"Defender")+COUNTIFS({CH("E")},F{r},{CH("F")},"Challenger"))', fmt="0")
    cell(ld, f"I{r}", f'=IF(F{r}="","",MATCH(F{r},$C$7:$C$26,0)-E{r})', fmt='"▲"0;"▼"0;"–"')
ld.conditional_formatting.add("I7:I26", FormulaRule(formula=['AND(ISNUMBER(I7),I7>0)'], font=Font(bold=True, color=GREEN)))
ld.conditional_formatting.add("I7:I26", FormulaRule(formula=['AND(ISNUMBER(I7),I7<0)'], font=Font(bold=True, color=RED)))
ld["C28"] = "List players from the top with no gaps. Names must match exactly on the Challenges tab."
ld["C28"].font = note_font
ld.print_area = "B2:I26"

# ------------------------------------------------------------------ Challenges
ch = wb.create_sheet("Challenges")
title(ch, "Challenges", "Log each ladder challenge in order", 8)
header(ch, 6, 2, ["No.", "Date", "Challenger", "Defender", "Winner", "Score", "What happened"], 30)
for c, w in zip("BCDEFGH", [6, 13, 22, 22, 13, 13, 34]):
    ch.column_dimensions[c].width = w
for c in "JKL":
    ch.column_dimensions[c].hidden = True
dv_player = DataValidation(type="list", formula1="Ladder!$C$7:$C$26", allow_blank=True)
dv_winner = DataValidation(type="list", formula1='"Challenger,Defender"', allow_blank=True)
ch.add_data_validation(dv_player)
ch.add_data_validation(dv_winner)
dv_player.add(f"D7:E{6 + CHALLENGES}")
dv_winner.add(f"F7:F{6 + CHALLENGES}")
sample_challenges = [
    (date(2026, 4, 4), "Sofia", "Priya", "Challenger", "6-4 6-3"),
    (date(2026, 4, 11), "Mateo", "Leah", "Defender", "4-6 6-7"),
    (date(2026, 4, 18), "Oscar", "Jonas", "Challenger", "7-5 6-4"),
    (date(2026, 4, 25), "Marcus", "Olivia", "Challenger", "6-7 6-3 10-8"),
]
for t in range(1, CHALLENGES + 1):
    r = 6 + t
    prev = col(2 + t)  # engine column holding positions before challenge t
    n = ch.cell(row=r, column=2, value=t)
    n.font, n.alignment, n.border = label_font, center, box
    for c in "CDEFG":
        cell(ch, f"{c}{r}", None, "input", "d mmm yyyy" if c == "C" else None, left if c in "DE" else center)
    if t <= len(sample_challenges):
        for c, v in zip("CDEFG", sample_challenges[t - 1]):
            ch[f"{c}{r}"] = v
    ch[f"J{r}"] = f"=IFERROR(INDEX(Engine!${prev}$7:${prev}$26,MATCH(D{r},Engine!$B$7:$B$26,0)),0)"
    ch[f"K{r}"] = f"=IFERROR(INDEX(Engine!${prev}$7:${prev}$26,MATCH(E{r},Engine!$B$7:$B$26,0)),0)"
    ch[f"L{r}"] = f'=IF(AND(D{r}<>"",E{r}<>"",D{r}<>E{r},F{r}="Challenger",J{r}>0,K{r}>0,J{r}>K{r}),1,0)'
    cell(ch, f"H{r}", f'=IF(OR(D{r}="",E{r}="",F{r}=""),"",IF(L{r}=1,D{r}&" moves up to #"&K{r},'
                      f'IF(F{r}="Challenger",D{r}&" wins · no change",E{r}&" stays at #"&K{r})))', align=left)
ch.freeze_panes = "B7"

# ------------------------------------------------------------------ Engine (hidden ladder history)
en = wb.create_sheet("Engine")
en["B2"] = "Ladder positions after each challenge. Calculated automatically; please don't edit."
en["B6"] = "Player"
for i in range(LADDER_MAX):
    r = 7 + i
    en[f"B{r}"] = f'=IF(Ladder!$C{r}="","",Ladder!$C{r})'
    en[f"C{r}"] = f'=IF($B{r}="","",ROW()-6)'
    for t in range(1, CHALLENGES + 1):
        c, p, lr = col(3 + t), col(2 + t), 6 + t
        en[f"{c}{r}"] = (f'=IF($B{r}="","",IF(Challenges!$L${lr}=1,IF($B{r}=Challenges!$D${lr},Challenges!$K${lr},'
                         f'IF(AND({p}{r}>=Challenges!$K${lr},{p}{r}<Challenges!$J${lr}),{p}{r}+1,{p}{r})),{p}{r}))')
en.sheet_state = "hidden"

for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = GREEN if sheet.title in ("Setup", "Schedule", "Standings") else RED
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
