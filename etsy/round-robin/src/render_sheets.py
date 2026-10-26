"""Renders each workbook tab to a PNG (for listing images). Run: python3 render_sheets.py <workbook.xlsx> <out_dir>"""
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageChops

src, out = Path(sys.argv[1]), Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)
AREAS = {"Setup": "B2:E31", "Schedule": "B2:H33", "Standings": "B2:K14", "Results Grid": "B2:I14",
         "Ladder": "B2:I17", "Challenges": "B2:H14"}
for sheet, area in AREAS.items():
    wb = load_workbook(src)
    for ws in wb.worksheets:
        ws.sheet_state = "visible" if ws.title == sheet else "hidden"
        ws.print_area = None  # LibreOffice exports every sheet that has a print area, hidden or not
    ws = wb[sheet]
    ws.print_area = area
    ws.print_title_rows = None
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    wb.active = wb.sheetnames.index(sheet)
    slug = sheet.lower().replace(" ", "-")
    tmp = out / f"{slug}.xlsx"
    wb.save(tmp)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(tmp)],
                   check=True, capture_output=True, timeout=180)
    subprocess.run(["pdftoppm", "-r", "220", "-png", "-singlefile", str(out / f"{slug}.pdf"), str(out / slug)], check=True)
    im = Image.open(out / f"{slug}.png").convert("RGB")
    bbox = ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()
    pad = 30
    im.crop((max(bbox[0] - pad, 0), max(bbox[1] - pad, 0), bbox[2] + pad, bbox[3] + pad)).save(out / f"{slug}.png")
    tmp.unlink()
    (out / f"{slug}.pdf").unlink()
    print(slug)
